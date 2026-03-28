import os
import re
import time
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook


# =========================
# PATHS
# =========================
PROJECT_ROOT = Path(__file__).resolve().parent.parent
ENV_PATH = PROJECT_ROOT / ".env"
DATA_INPUT_DIR = PROJECT_ROOT / "data" / "input"

load_dotenv(ENV_PATH)


# =========================
# CONFIG
# =========================
WB_TOKEN = os.getenv("WB_TOKEN")
PHOTOS_EXCEL_FILE = os.getenv("PHOTOS_EXCEL_FILE", "photos.xlsx")
PHOTOS_START_ROW = int(os.getenv("PHOTOS_START_ROW", "2"))
PHOTOS_SHEET_NAME = os.getenv("PHOTOS_SHEET_NAME")
PHOTOS_SLEEP = float(os.getenv("PHOTOS_SLEEP", "1.0"))
PHOTOS_DRY_RUN = os.getenv("PHOTOS_DRY_RUN", "true").strip().lower() == "true"

if not WB_TOKEN:
    raise ValueError("WB_TOKEN не найден в .env")


# =========================
# API
# =========================
URL = "https://content-api.wildberries.ru/content/v3/media/save"

HEADERS = {
    "Authorization": WB_TOKEN,
    "Content-Type": "application/json",
}


# =========================
# HELPERS
# =========================
def parse_photo_urls(raw_value):
    if raw_value is None:
        return []

    text = str(raw_value).strip()
    if not text:
        return []

    parts = re.split(r"[\n,;]+", text)
    return [p.strip() for p in parts if p.strip()]


def send_photos(session, nm_id, photo_urls):
    payload = {
        "nmId": nm_id,
        "data": photo_urls,
    }

    for attempt in range(1, 4):
        try:
            resp = session.post(URL, headers=HEADERS, json=payload, timeout=120)

            print(f"[nmId {nm_id}] STATUS: {resp.status_code}")

            if resp.status_code == 200:
                print(f"[nmId {nm_id}] OK ({len(photo_urls)} фото)")
                return True
            else:
                print(resp.text)
                print(f"[nmId {nm_id}] Ошибка, попытка {attempt}/3")
                time.sleep(2 * attempt)

        except requests.exceptions.RequestException as e:
            print(f"[nmId {nm_id}] Сетевая ошибка, попытка {attempt}/3: {e}")
            time.sleep(2 * attempt)

    return False


# =========================
# MAIN
# =========================
def main():
    excel_path = DATA_INPUT_DIR / PHOTOS_EXCEL_FILE

    if not excel_path.exists():
        raise FileNotFoundError(f"Не найден файл: {excel_path}")

    wb = load_workbook(excel_path, data_only=True)

    if PHOTOS_SHEET_NAME:
        if PHOTOS_SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"Лист '{PHOTOS_SHEET_NAME}' не найден. Доступные листы: {wb.sheetnames}"
            )
        ws = wb[PHOTOS_SHEET_NAME]
    else:
        ws = wb[wb.sheetnames[0]]

    session = requests.Session()

    total_rows = 0
    success_count = 0
    error_count = 0
    skipped_count = 0

    for row_num in range(PHOTOS_START_ROW, ws.max_row + 1):
        nm_id_raw = ws.cell(row=row_num, column=1).value
        photos_raw = ws.cell(row=row_num, column=2).value

        if nm_id_raw is None and photos_raw is None:
            continue

        total_rows += 1

        if nm_id_raw is None:
            print(f"[Строка {row_num}] Пропуск: пустой nmId")
            skipped_count += 1
            continue

        try:
            nm_id = int(str(nm_id_raw).strip())
        except ValueError:
            print(f"[Строка {row_num}] Пропуск: некорректный nmId = {nm_id_raw}")
            skipped_count += 1
            continue

        photo_urls = parse_photo_urls(photos_raw)

        if not photo_urls:
            print(f"[Строка {row_num}] Пропуск: нет ссылок для nmId {nm_id}")
            skipped_count += 1
            continue

        print(f"\n[Строка {row_num}] nmId: {nm_id}")
        print(f"[Строка {row_num}] Ссылок: {len(photo_urls)}")

        if PHOTOS_DRY_RUN:
            print(f"[Строка {row_num}] PHOTOS_DRY_RUN=True -> фото не отправлялись")
            success_count += 1
            continue

        ok = send_photos(session, nm_id, photo_urls)

        if ok:
            success_count += 1
        else:
            error_count += 1
            print(f"[Строка {row_num}] Не удалось загрузить фото для nmId {nm_id}")

        time.sleep(PHOTOS_SLEEP)

    print("\n==== ИТОГ ====")
    print(f"Обработано строк: {total_rows}")
    print(f"Успешно: {success_count}")
    print(f"С ошибками: {error_count}")
    print(f"Пропущено: {skipped_count}")


if __name__ == "__main__":
    main()